#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
spreadsheet.py — 表格生成 / 解析（本地兜底，零依赖）

把「原文 / 摘要」中的结构化数据转为电子表格可消费的形态，是 `spreadsheet`
协同能力（摘要→表格）**缺失时的本地兜底**：

  - 输入：纯文本，或 `structured_summary.py --format json` 的产物（`--from-json`）；
  - 抽取：问答对 / 列表项 / 定义说明 / 表格行（复用 structured_summary 的启发式）；
  - 输出：
      * csv  —— 规整的 CSV 文本（可直接粘贴进 Excel / 飞书多维表 / Numbers）；
      * html —— 自包含 HTML，每类数据一段 <table>，浏览器打开即「直接渲染」；
      * md   —— 每类一段 markdown 表格。

纯标准库实现（零依赖、跨平台）。默认输出 CSV / 自包含 HTML / markdown 表，
已覆盖绝大多数「表格」下游需求；**可选** `--format xlsx` 经 openpyxl 生成真正的
.xlsx 工作簿（每类数据一个 sheet），未安装 openpyxl 时优雅降级并提示安装命令，
不影响零依赖默认路径。
"""

import sys
import os
import re
import csv
import json
import io
import argparse

# 复用本技能自带的启发式抽取（同目录、零依赖，非外部技能）
try:
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from structured_summary import extract_items
except Exception:  # pragma: no cover
    extract_items = None


def read_text(path):
    if path == "-":
        if hasattr(sys.stdin, "reconfigure"):
            try:
                sys.stdin.reconfigure(encoding="utf-8")
            except Exception:
                pass
        return sys.stdin.read()
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


# ---------------------------------------------------------------------------
# 把抽取条目归一成 (类别标签, [字段...]) 列表
# ---------------------------------------------------------------------------
TYPE_LABEL = {
    "qa": "问答对",
    "list": "列表项",
    "definition": "定义/说明",
    "table": "表格",
}


def normalize(items):
    """返回 {type: [ [field,...], ... ]} 并按类别顺序组织。"""
    grouped = {}
    header_seen = False  # 表格首行视为表头
    for it in items:
        t = it["type"]
        if t == "qa":
            row = [it.get("question", ""), it.get("answer", "")]
        elif t == "list":
            row = [it.get("content", "")]
        elif t == "definition":
            row = [it.get("term", ""), it.get("desc", "")]
        elif t == "table":
            row = list(it.get("cells", []))
            if not header_seen:
                # 首条表格行作为表头（不计入数据），其后为数据
                header_seen = True
                grouped.setdefault("_table_header", row)
                continue
        else:
            row = [it.get("content", "")]
        grouped.setdefault(t, []).append(row)
    return grouped


# ---------------------------------------------------------------------------
# CSV 渲染（统一列式：类型 + 最多 3 个值字段）
# ---------------------------------------------------------------------------
def render_csv(grouped):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["类型", "值1", "值2", "值3"])
    order = ["qa", "list", "definition", "table"]
    for t in order:
        rows = grouped.get(t)
        if not rows:
            continue
        for r in rows:
            vals = list(r)
            if len(vals) > 3:
                vals = vals[:2] + [" | ".join(vals[2:])]
            vals = (vals + ["", "", ""])[:3]
            w.writerow([TYPE_LABEL.get(t, t)] + vals)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# HTML 渲染（每类一段 <table>，直接渲染）
# ---------------------------------------------------------------------------
def _esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def render_html(grouped):
    parts = ['<!DOCTYPE html>', '<html lang="zh-CN">',
             '<head><meta charset="utf-8">',
             '<title>结构化数据表格</title>',
             '<style>body{font-family:system-ui,sans-serif;margin:24px;color:#222}'
             'h2{margin:1.4em 0 .4em;font-size:1.05em}'
             'table{border-collapse:collapse;margin-bottom:8px}'
             'th,td{border:1px solid #cfd8dc;padding:6px 10px;font-size:13px;text-align:left}'
             'th{background:#f5f7fa}caption{caption-side:top;font-weight:600}</style>',
             '</head><body>']
    parts.append("<h1>结构化数据表格</h1>")
    order = ["qa", "list", "definition", "table"]
    for t in order:
        rows = grouped.get(t)
        if not rows:
            continue
        parts.append("<h2>%s（%d）</h2>" % (_esc(TYPE_LABEL.get(t, t)), len(rows)))
        parts.append("<table>")
        if t == "table" and grouped.get("_table_header"):
            parts.append("<tr>" + "".join(
                "<th>%s</th>" % _esc(c) for c in grouped["_table_header"]) + "</tr>")
        for r in rows:
            parts.append("<tr>" + "".join("<td>%s</td>" % _esc(c) for c in r) + "</tr>")
        parts.append("</table>")
    parts.append("</body></html>")
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Markdown 渲染
# ---------------------------------------------------------------------------
def render_md(grouped):
    lines = ["# 结构化数据表格", ""]
    order = ["qa", "list", "definition", "table"]
    for t in order:
        rows = grouped.get(t)
        if not rows:
            continue
        lines.append("## %s（%d）" % (TYPE_LABEL.get(t, t), len(rows)))
        lines.append("")
        if t == "table" and grouped.get("_table_header"):
            hdr = grouped["_table_header"]
            lines.append("| " + " | ".join(hdr) + " |")
            lines.append("| " + " | ".join(["---"] * len(hdr)) + " |")
            for r in rows:
                lines.append("| " + " | ".join(r) + " |")
        elif t == "qa":
            lines.append("| 问题 | 回答 |")
            lines.append("| --- | --- |")
            for r in rows:
                lines.append("| %s | %s |" % (r[0], r[1] if len(r) > 1 else ""))
        elif t == "definition":
            lines.append("| 术语 | 释义 |")
            lines.append("| --- | --- |")
            for r in rows:
                lines.append("| %s | %s |" % (r[0], r[1] if len(r) > 1 else ""))
        else:  # list
            lines.append("| 内容 |")
            lines.append("| --- |")
            for r in rows:
                lines.append("| %s |" % r[0])
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


# ---------------------------------------------------------------------------
# XLSX 渲染（可选；经 openpyxl，未安装则优雅降级）
# ---------------------------------------------------------------------------
def render_xlsx(grouped, out_path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception:
        sys.stderr.write("错误：生成 .xlsx 需要 openpyxl。请先安装：\n"
                         "  pip install openpyxl\n然后重试 --format xlsx。\n")
        sys.exit(2)

    wb = Workbook()
    wb.remove(wb.active)  # 移除默认空白 sheet
    head_fill = PatternFill("solid", fgColor="F5F7FA")
    head_font = Font(bold=True)
    wrap = Alignment(vertical="top", wrap_text=True)

    specs = [
        ("qa", "问答对", ["问题", "回答"]),
        ("list", "列表项", ["内容"]),
        ("definition", "定义说明", ["术语", "释义"]),
        ("table", "表格", None),  # 表头取 _table_header
    ]
    for t, label, headers in specs:
        rows = grouped.get(t)
        if not rows:
            continue
        ws = wb.create_sheet(title=label[:31])
        hdr = grouped.get("_table_header") if t == "table" else headers
        ws.append(list(hdr))
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = wrap
        for r in rows:
            ws.append([c if c is not None else "" for c in r])
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 12), 60)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap

    if not wb.sheetnames:
        ws = wb.create_sheet(title="无数据")
        ws.append(["未从输入抽取到结构化条目"])

    wb.save(out_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="表格生成/解析（本地兜底，零依赖；可选 xlsx）")
    ap.add_argument("input", help="文本文件，或 structured_summary 的 JSON（配合 --from-json）；- 读 stdin")
    ap.add_argument("--from-json", action="store_true",
                    help="input 为 structured_summary.py --format json 的产物，直接读 items")
    ap.add_argument("--mode", default="auto",
                    choices=["auto", "qa", "list", "definition", "table"],
                    help="抽取类型子集（默认 auto=全部）")
    ap.add_argument("--format", default="csv", choices=["csv", "html", "md", "xlsx"])
    ap.add_argument("--out", default=None, help="输出文件路径（html/csv/xlsx 建议落盘）")
    args = ap.parse_args()

    if args.from_json:
        data = json.loads(read_text(args.input))
        items = data.get("items", []) if isinstance(data, dict) else []
        stats = data.get("stats", {})
    else:
        if extract_items is None:
            sys.stderr.write("错误：未能加载 structured_summary.extract_items\n")
            sys.exit(2)
        r = extract_items(read_text(args.input), mode=args.mode)
        items = r["items"]
        stats = r["stats"]

    grouped = normalize(items)
    if args.format == "xlsx":
        out_path = args.out or "structured_data.xlsx"
        render_xlsx(grouped, out_path)
        sys.stderr.write("已写入：%s（条目=%d）\n"
                         % (out_path, sum(len(v) for k, v in grouped.items() if k != "_table_header")))
        return
    elif args.format == "csv":
        out = render_csv(grouped)
    elif args.format == "html":
        out = render_html(grouped)
    else:
        out = render_md(grouped)

    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            f.write(out)
        sys.stderr.write("已写入：%s（条目=%d）\n" % (args.out, sum(len(v) for k, v in grouped.items() if k != "_table_header")))
    else:
        sys.stdout.write(out)


if __name__ == "__main__":
    main()
